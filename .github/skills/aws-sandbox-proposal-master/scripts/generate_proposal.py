#!/usr/bin/env python3
"""
Standalone AWS Sandbox Proposal DOCX Generator (Template-Based).

Binds JSON context data onto the standard "Sandbox Innovation Plan Template_v2.docx"
using docxtpl (Jinja2 for Word). Markdown content in the JSON is auto-converted to
RichText so that bold, italic, lists, headers, etc. render correctly in the final DOCX.

Usage:
    python generate_proposal.py <context.json> [--output proposal.docx] [--template path/to/template.docx]

Dependencies:
    pip install docxtpl python-docx
"""
import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple, Union

from docxtpl import DocxTemplate, InlineImage, RichText
from docx.shared import Inches


# ---------------------------------------------------------------------------
# Default template path (relative to this script)
# ---------------------------------------------------------------------------
_SCRIPT_DIR = Path(__file__).resolve().parent
_DEFAULT_TEMPLATE = _SCRIPT_DIR.parent / 'assets' / 'Sandbox Innovation Plan Template_v2.docx'


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------
_CALCULATOR_LINK_PATTERN = re.compile(r'^https://calculator\.aws/#/estimate\?id=[A-Za-z0-9_-]+$')


def _normalized_name_set(values: Sequence[str]) -> set:
    return {v.strip().lower() for v in values if isinstance(v, str) and v.strip()}


def _extract_service_names(service_list: Any) -> List[str]:
    names: List[str] = []
    if not isinstance(service_list, list):
        return names
    for item in service_list:
        if isinstance(item, dict):
            name = item.get('service_name', '')
            if isinstance(name, str) and name.strip():
                names.append(name.strip())
    return names


def validate_context(data: dict, context_dir: Optional[Path] = None) -> List[str]:
    """Return list of missing required field paths."""
    missing: List[str] = []
    sandbox = data.get('sandbox')
    if not isinstance(sandbox, dict):
        return ['sandbox (root object)']

    basics = sandbox.get('basics', {})
    for key in ('title', 'partner', 'pdm', 'sa'):
        if not basics.get(key):
            missing.append(f'sandbox.basics.{key}')
    contact = basics.get('contact', {})
    for key in ('name', 'title', 'email'):
        if not contact.get(key):
            missing.append(f'sandbox.basics.contact.{key}')

    details = sandbox.get('details', {})
    for key in ('summary', 'features', 'solution_type', 'customer_type', 'pain_point'):
        if not details.get(key):
            missing.append(f'sandbox.details.{key}')

    business = sandbox.get('business', {})
    for key in ('justification', 'aws_funding', 'labor_cost', 'total_cost',
                'start_date', 'end_date', 'release_date'):
        if not business.get(key):
            missing.append(f'sandbox.business.{key}')

    region = business.get('region', '')
    if not isinstance(region, str) or not region.strip():
        missing.append('sandbox.business.region')

    service_list = business.get('service_list')
    if not isinstance(service_list, list) or not service_list:
        missing.append('sandbox.business.service_list')
    else:
        for i, service in enumerate(service_list):
            if not isinstance(service, dict):
                missing.append(f'sandbox.business.service_list[{i}]')
                continue
            if not service.get('service_name'):
                missing.append(f'sandbox.business.service_list[{i}].service_name')

    # calculator_link is MANDATORY and must be a valid AWS Calculator share URL.
    # An empty string, missing key, or placeholder is a hard failure — the proposal
    # must NOT be generated before Step 3B (AWS Pricing Calculator) is completed.
    calculator_link = business.get('calculator_link', '')
    if not isinstance(calculator_link, str) or not calculator_link.strip():
        missing.append(
            'sandbox.business.calculator_link (REQUIRED — complete Step 3B '
            'in the AWS Pricing Calculator before generating this document)'
        )
    elif not _CALCULATOR_LINK_PATTERN.match(calculator_link.strip()):
        missing.append(
            'sandbox.business.calculator_link (invalid format — must be '
            'https://calculator.aws/#/estimate?id=<id>)'
        )

    plan = sandbox.get('plan', {})
    if not plan.get('total_mandays'):
        missing.append('sandbox.plan.total_mandays')
    phases = plan.get('total_phases', [])
    if not phases:
        missing.append('sandbox.plan.total_phases')
    else:
        for i, phase in enumerate(phases):
            for key in ('activity', 'description', 'mandays', 'delivery_date'):
                if not phase.get(key):
                    missing.append(f'sandbox.plan.total_phases[{i}].{key}')

    arch = sandbox.get('architecture', {})
    if not arch.get('description'):
        missing.append('sandbox.architecture.description')

    diagram_value = arch.get('diagram', '')
    if not isinstance(diagram_value, str) or not diagram_value.strip():
        missing.append('sandbox.architecture.diagram')
    else:
        dp = Path(diagram_value)
        if not dp.is_absolute() and context_dir is not None:
            dp = context_dir / dp
        if not dp.exists():
            missing.append(f'sandbox.architecture.diagram (file not found: {diagram_value})')

    expected_service_names = _extract_service_names(service_list)
    expected_services = _normalized_name_set(expected_service_names)

    diagram_services = _normalized_name_set(arch.get('applied_services', []))
    if diagram_services and diagram_services != expected_services:
        missing.append('sandbox.architecture.applied_services (mismatch vs sandbox.business.service_list)')

    calculator_services = _normalized_name_set(business.get('applied_services', []))
    if calculator_services and calculator_services != expected_services:
        missing.append('sandbox.business.applied_services (mismatch vs sandbox.business.service_list)')

    return missing


# ---------------------------------------------------------------------------
# Markdown → RichText conversion  (minimal standalone DocxTemplator)
# ---------------------------------------------------------------------------
def _is_markdown(text: str) -> bool:
    """Detect if a string contains Markdown formatting."""
    if not isinstance(text, str) or len(text) < 3:
        return False
    if re.match(r'^[a-zA-Z0-9_\.]+$', text):
        return False
    if len(text.strip()) < 10 and re.match(r'^[a-zA-Z0-9_\s\-\.,:;]+$', text):
        return False
    patterns = [
        r'^#{1,6}\s+', r'\*\*[^*]+\*\*', r'__[^_]+__',
        r'^\s*[-*]\s+\w', r'^\s*\d+\.\s+\w', r'\[.+?\]\(.+?\)',
        r'`[^`]+`', r'```', r'^>\s+\w',
    ]
    return any(re.search(p, text, re.MULTILINE) for p in patterns)


def _parse_inline(rt: RichText, text: str, *, bold: bool = False, italic: bool = False):
    """Parse inline bold / italic / code / link formatting into a RichText."""
    pattern = re.compile(
        r'(\*\*[^*]+?\*\*)'
        r'|(__[^_]+?__)'
        r'|(?<!\*)\*(?!\*)([^*]+?)\*(?!\*)'
        r'|(`[^`]+?`)'
        r'|(\[.+?\]\(.+?\))'
    )
    last = 0
    for m in pattern.finditer(text):
        # plain text before this match
        if m.start() > last:
            rt.add(text[last:m.start()], bold=bold, italic=italic)
        seg = m.group(0)
        if seg.startswith('**') and seg.endswith('**'):
            rt.add(seg[2:-2], bold=True, italic=italic)
        elif seg.startswith('__') and seg.endswith('__'):
            rt.add(seg[2:-2], bold=True, italic=italic)
        elif seg.startswith('`') and seg.endswith('`'):
            rt.add(seg[1:-1], font='Consolas', size=20, color='C7254E')
        elif seg.startswith('['):
            lm = re.match(r'\[(.+?)\]\((.+?)\)', seg)
            if lm:
                rt.add(lm.group(1), color='0563C1', underline=True, bold=bold, italic=italic)
        else:
            # italic *text*
            inner = m.group(3)
            if inner:
                rt.add(inner, bold=bold, italic=True)
        last = m.end()
    if last < len(text):
        rt.add(text[last:], bold=bold, italic=italic)


def markdown_to_richtext(md_text: str) -> RichText:
    """Convert a Markdown string to a docxtpl RichText object."""
    if not md_text:
        return RichText('')
    rt = RichText()
    lines = md_text.strip().split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            rt.add('\n')
            i += 1
            continue

        # Code blocks
        if stripped.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                rt.add(lines[i], font='Consolas', size=20, color='2E3440')
                rt.add('\n')
                i += 1
            if i < len(lines):
                i += 1  # skip closing ```
            rt.add('\n')
            continue

        # Blockquotes
        qm = re.match(r'^>\s*(.+)$', stripped)
        if qm:
            rt.add('  ┃ ', color='808080')
            _parse_inline(rt, qm.group(1))
            if i < len(lines) - 1:
                rt.add('\n')
            i += 1
            continue

        # Headers
        hm = re.match(r'^(#{1,6})\s+(.+)$', stripped)
        if hm:
            level = len(hm.group(1))
            sizes = {1: 28, 2: 24, 3: 20}
            rt.add(hm.group(2), bold=True, size=sizes.get(level, 18))
            if i < len(lines) - 1:
                rt.add('\n')
            i += 1
            continue

        # Unordered list
        lm = re.match(r'^[-*]\s+(.+)$', stripped)
        if lm:
            rt.add('  • ')
            _parse_inline(rt, lm.group(1))
            if i < len(lines) - 1:
                rt.add('\n')
            i += 1
            continue

        # Ordered list
        om = re.match(r'^(\d+)\.\s+(.+)$', stripped)
        if om:
            rt.add(f'  {om.group(1)}. ')
            _parse_inline(rt, om.group(2))
            if i < len(lines) - 1:
                rt.add('\n')
            i += 1
            continue

        # Inline code line (has backticks)
        if '`' in stripped:
            parts = re.split(r'(`[^`]+?`)', stripped)
            for part in parts:
                if part.startswith('`') and part.endswith('`'):
                    rt.add(part[1:-1], font='Consolas', size=20, color='C7254E')
                else:
                    _parse_inline(rt, part)
            if i < len(lines) - 1:
                rt.add('\n')
            i += 1
            continue

        # Regular paragraph
        _parse_inline(rt, stripped)
        if i < len(lines) - 1:
            rt.add('\n')
        i += 1

    return rt


# Keys whose values must stay as plain strings in the rendered context.
# The template stores these placeholders inside split-run CJK-formatted table
# cells (multiple <w:r> elements with different rPr nodes). docxtpl cannot
# inject a RichText object into a split-run placeholder — only plain strings
# survive the merge correctly.  Strip Markdown syntax instead of converting.
_PLAIN_TEXT_KEYS: frozenset = frozenset({
    'summary', 'description', 'features', 'pain_point',
    'justification', 'additional_info',
})


def markdown_to_plain(text: str) -> str:
    """Strip Markdown syntax from *text* and return clean plain text."""
    if not isinstance(text, str):
        return text
    # Remove fenced code blocks entirely (keep content)
    text = re.sub(r'```[^\n]*\n(.*?)```', r'\1', text, flags=re.DOTALL)
    text = re.sub(r'```', '', text)
    # Remove ATX headers (## Title → Title)
    text = re.sub(r'^#{1,6}\s+', '', text, flags=re.MULTILINE)
    # Remove bold / italic markers
    text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
    text = re.sub(r'__([^_]+)__', r'\1', text)
    text = re.sub(r'(?<!\*)\*(?!\*)([^*]+?)\*(?!\*)', r'\1', text)
    # Remove inline code
    text = re.sub(r'`([^`]+)`', r'\1', text)
    # Remove links — keep display text
    text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)
    # Remove blockquote markers
    text = re.sub(r'^>\s*', '', text, flags=re.MULTILINE)
    # Normalise unordered list markers to bullet
    text = re.sub(r'^\s*[-*]\s+', '• ', text, flags=re.MULTILINE)
    return text.strip()


def auto_convert_markdown(data: Any, _current_key: str = '') -> Any:
    """Recursively walk a dict/list and convert Markdown strings to RichText.

    Fields listed in *_PLAIN_TEXT_KEYS* are stripped to plain text instead of
    being converted to RichText objects, because the DOCX template renders those
    placeholders inside split-run CJK-formatted cells that cannot accept RichText.
    """
    if isinstance(data, dict):
        return {k: auto_convert_markdown(v, _current_key=k) for k, v in data.items()}
    if isinstance(data, list):
        return [auto_convert_markdown(item, _current_key=_current_key) for item in data]
    if isinstance(data, str):
        if _current_key in _PLAIN_TEXT_KEYS:
            return markdown_to_plain(data)
        if _is_markdown(data):
            return markdown_to_richtext(data)
    return data


# ---------------------------------------------------------------------------
# Context reshaping — raw JSON → template-ready context
# ---------------------------------------------------------------------------
def build_render_context(
    payload: Dict[str, Any],
    diagram_path: Optional[Path],
) -> Dict[str, Any]:
    """Reshape the raw JSON into the structure expected by the DOCX template.

    The template uses Jinja2 placeholders like ``{{sandbox.details.summary}}``,
    ``{{sandbox.business.labor_cost}}``, and ``{%tr for phase in
    sandbox.plan.total_phases%}`` loops.
    """
    sandbox = payload['sandbox']
    basics = sandbox['basics']
    details = sandbox['details']
    business = sandbox['business']
    plan = sandbox['plan']
    architecture = sandbox['architecture']

    contact = dict(basics.get('contact', {}))
    contact['pdm'] = basics.get('pdm', contact.get('pdm', ''))
    contact['pm'] = basics.get('sa', contact.get('pm', ''))

    return {
        'sandbox': {
            'proposal': {'title': basics.get('title', '')},
            'basics': {
                'title': basics.get('title', ''),
                'partner': basics.get('partner', ''),
                'contact': contact,
                'pdm': basics.get('pdm', ''),
                'sa': basics.get('sa', ''),
            },
            'details': details,
            'business': business,
            'plan': plan,
            'architecture': {
                **architecture,
                'diagram': diagram_path,  # will be replaced with InlineImage
            },
        }
    }


# ---------------------------------------------------------------------------
# Main generator
# ---------------------------------------------------------------------------
def generate_proposal(
    context: dict,
    output_path: Path,
    template_path: Optional[Path] = None,
    context_dir: Optional[Path] = None,
) -> str:
    """Generate a Sandbox Innovation Plan DOCX by binding context onto the template.

    Args:
        context: The full sandbox context dictionary.
        output_path: Where to save the generated DOCX.
        template_path: Path to the DOCX template. Defaults to the bundled asset.
        context_dir: Base directory for resolving relative paths (e.g. diagram).
    """
    if template_path is None:
        template_path = _DEFAULT_TEMPLATE
    if context_dir is None:
        context_dir = output_path.parent

    # --- validate ---
    missing = validate_context(context, context_dir=context_dir)
    if missing:
        print('ERROR: Missing required fields:\n  ' + '\n  '.join(missing),
              file=sys.stderr)
        sys.exit(1)

    if not template_path.exists():
        print(f'ERROR: Template not found: {template_path}', file=sys.stderr)
        sys.exit(1)

    # --- resolve diagram ---
    sandbox = context['sandbox']
    diagram_path: Optional[Path] = None
    diagram_str = sandbox.get('architecture', {}).get('diagram', '')
    if diagram_str:
        dp = Path(diagram_str)
        if not dp.is_absolute():
            dp = context_dir / dp
        diagram_path = dp

    # --- build render context ---
    render_ctx = build_render_context(context, diagram_path)

    # --- load template & prepare InlineImage ---
    doc = DocxTemplate(str(template_path))

    if diagram_path:
        # Fit image within printable area: max 6 in wide, max 7 in tall.
        # Read PNG dimensions (bytes 16-24 of IHDR chunk) to compute aspect ratio.
        max_w = Inches(6.0)
        max_h = Inches(7.0)
        try:
            import struct
            with open(str(diagram_path), 'rb') as _f:
                _f.read(16)  # sig (8) + IHDR length (4) + 'IHDR' (4)
                _pw = struct.unpack('>I', _f.read(4))[0]
                _ph = struct.unpack('>I', _f.read(4))[0]
            _aspect = _ph / _pw if _pw else 1.0
            _h_at_max_w = max_w * _aspect
            if _h_at_max_w <= max_h:
                _img_kwargs = {'width': max_w}
            else:
                _img_kwargs = {'height': max_h}
        except Exception:
            _img_kwargs = {'width': max_w}
        render_ctx['sandbox']['architecture']['diagram'] = InlineImage(
            doc, str(diagram_path), **_img_kwargs,
        )

    # --- auto-convert Markdown fields to RichText ---
    render_ctx = auto_convert_markdown(render_ctx)

    # --- render & save ---
    doc.render(render_ctx)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return str(output_path)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(
        description='Generate AWS Sandbox Innovation Plan DOCX from JSON context.'
    )
    parser.add_argument('context_json', help='Path to the sandbox context JSON file')
    parser.add_argument('--output', '-o', default=None,
                        help='Output DOCX path (default: same name as JSON with .docx)')
    parser.add_argument('--template', '-t', default=None,
                        help='Path to the DOCX template (default: bundled asset)')
    args = parser.parse_args()

    json_path = Path(args.context_json)
    if not json_path.exists():
        print(f'ERROR: Context file not found: {json_path}', file=sys.stderr)
        sys.exit(1)

    context = json.loads(json_path.read_text(encoding='utf-8'))
    if 'document_data' in context and isinstance(context['document_data'], dict):
        context = context['document_data']

    output = Path(args.output) if args.output else json_path.with_suffix('.docx')
    template = Path(args.template) if args.template else None

    result = generate_proposal(context, output, template_path=template,
                               context_dir=json_path.parent)
    print(f'Proposal generated: {result}')


if __name__ == '__main__':
    main()
